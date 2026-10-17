"""Carbon Emissions Reporting Platform - Flask + SQLite backend."""
import os, random, sqlite3, datetime as dt, calendar
import openpyxl
from flask import Flask, g, jsonify, request, send_from_directory

DB_PATH = os.environ.get("DB_PATH", "ghg.db")
XLSX = os.path.join(os.path.dirname(__file__), "data", "GHG_Sheet_.xlsx")
app = Flask(__name__, static_folder="static")

SCHEMA = """
CREATE TABLE IF NOT EXISTS emission_factors(
  id INTEGER PRIMARY KEY, scope INTEGER NOT NULL, activity TEXT NOT NULL, unit TEXT NOT NULL,
  kgco2e_per_unit REAL NOT NULL, source TEXT, valid_from TEXT NOT NULL, valid_to TEXT,  -- NULL = still valid
  UNIQUE(scope, activity, valid_from));
CREATE TABLE IF NOT EXISTS emission_records(
  id INTEGER PRIMARY KEY, scope INTEGER NOT NULL, activity TEXT NOT NULL, section TEXT, plant TEXT,
  activity_date TEXT NOT NULL, quantity REAL NOT NULL, unit TEXT NOT NULL,
  factor_id INTEGER NOT NULL REFERENCES emission_factors(id), factor_used REAL NOT NULL,
  calculated_kgco2e REAL NOT NULL, final_kgco2e REAL NOT NULL, is_overridden INTEGER DEFAULT 0,
  origin TEXT DEFAULT 'user', notes TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE INDEX IF NOT EXISTS ix_rec_date ON emission_records(activity_date, scope);
CREATE TABLE IF NOT EXISTS audit_log(
  id INTEGER PRIMARY KEY, record_id INTEGER NOT NULL REFERENCES emission_records(id),
  field TEXT, old_value REAL, new_value REAL, reason TEXT NOT NULL, changed_by TEXT,
  changed_at TEXT DEFAULT CURRENT_TIMESTAMP);
CREATE TABLE IF NOT EXISTS business_metrics(
  id INTEGER PRIMARY KEY, date TEXT NOT NULL, metric_name TEXT NOT NULL, value REAL NOT NULL,
  unit TEXT, aggregation TEXT DEFAULT 'sum', UNIQUE(date, metric_name));
"""

# ---------- db helpers ----------
def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH); g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d: d.close()

def rows(sql, args=()):
    return [dict(r) for r in db().execute(sql, args).fetchall()]

def find_factor(con, scope, activity, date):
    """Historical accuracy: the factor whose validity window contains the ACTIVITY date."""
    return con.execute(
        "SELECT * FROM emission_factors WHERE scope=? AND activity=? AND valid_from<=? "
        "AND (valid_to IS NULL OR valid_to>=?) ORDER BY valid_from DESC LIMIT 1",
        (scope, activity, date, date)).fetchone()

def create_record(con, scope, activity, qty, date, section=None, plant=None, notes=None, origin="user"):
    f = find_factor(con, scope, activity, date)
    if not f:
        raise ValueError(f"No emission factor valid on {date} for '{activity}' (scope {scope})")
    kg = qty * f["kgco2e_per_unit"]
    cur = con.execute(
        "INSERT INTO emission_records(scope,activity,section,plant,activity_date,quantity,unit,factor_id,"
        "factor_used,calculated_kgco2e,final_kgco2e,origin,notes) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (scope, activity, section, plant, date, qty, f["unit"], f["id"], f["kgco2e_per_unit"], kg, kg, origin, notes))
    return cur.lastrowid

# ---------- seeding from the provided spreadsheet ----------
def month_end(y, m): return f"{y}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}"
QMONTHS = {"Q1": (1, 2, 3), "Q2": (4, 5, 6), "Q3": (7, 8, 9), "Q4": (10, 11, 12)}

def init_db():
    con = sqlite3.connect(DB_PATH); con.row_factory = sqlite3.Row; con.executescript(SCHEMA)
    if con.execute("SELECT COUNT(*) FROM emission_factors").fetchone()[0]:
        con.close(); return
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    acts = {}   # (scope, activity) -> dict(unit, base_t, source, plant, entries[(section, quarter, qty)])
    for r in list(wb["Scope 1"].iter_rows(min_row=2, values_only=True)):
        sec, mat, unit, q, qty, fac, _, _, _, _, src, plant = r[:12]
        a = acts.setdefault((1, mat), dict(unit=unit, t=fac, src=src, plant=plant, e=[]))
        a["e"].append((sec, q, qty))
    for r in list(wb["Scope 2"].iter_rows(min_row=2, values_only=True)):
        sec, etype, supplier, unit, qty, fac = r[:6]
        src, plant, q = r[11], r[14], r[17]
        a = acts.setdefault((2, f"{etype} - {supplier}"), dict(unit=unit, t=fac, src=src, plant=plant, e=[]))
        a["e"].append((sec, q, qty))
    # versioned factors: 2023 (expired), 2024 = spreadsheet value (expired), 2025+ (current)
    for (scope, act), a in acts.items():
        for vf, vt, mult, src in (("2023-01-01", "2023-12-31", 1.04, "SAMPLE prior revision (illustrative)"),
                                  ("2024-01-01", "2024-12-31", 1.00, a["src"]),
                                  ("2025-01-01", None, 0.98, "SAMPLE current revision (illustrative)")):
            con.execute("INSERT INTO emission_factors(scope,activity,unit,kgco2e_per_unit,source,valid_from,valid_to)"
                        " VALUES(?,?,?,?,?,?,?)", (scope, act, a["unit"], a["t"] * 1000 * mult, src, vf, vt))
    rnd = random.Random(42)
    for (scope, act), a in acts.items():
        have = {}
        for sec, q, qty in a["e"]:                       # real sheet rows: quarter qty split evenly over 3 months
            for m in QMONTHS[q]:
                create_record(con, scope, act, qty / 3, month_end(2024, m), sec, a["plant"],
                              f"From {os.path.basename(XLSX)} ({q}, split evenly across months)", "sheet")
                have[(2024, m)] = 1
        base = sum(e[2] for e in a["e"]) / len(a["e"]) / 3     # average monthly quantity
        sec0 = a["e"][0][0]
        for y in (2024, 2025):                            # fill gaps + 2025 with clearly-labelled sample data
            for m in range(1, 13):
                if (y, m) in have: continue
                qty = base * rnd.uniform(0.92, 1.08) * (0.97 if y == 2025 else 1.0)
                create_record(con, scope, act, qty, month_end(y, m), sec0, a["plant"], "SAMPLE data (generated)", "sample")
    for y in (2024, 2025):
        for m in range(1, 13):
            con.execute("INSERT INTO business_metrics(date,metric_name,value,unit,aggregation) VALUES(?,?,?,?,?)",
                        (month_end(y, m), "Tons of Steel Produced", round(330000 * rnd.uniform(.95, 1.05)), "tonnes", "sum"))
            con.execute("INSERT INTO business_metrics(date,metric_name,value,unit,aggregation) VALUES(?,?,?,?,?)",
                        (month_end(y, m), "Number of Employees", round(12500 * rnd.uniform(.98, 1.02)), "people", "avg"))
    con.commit(); con.close()

# ---------- API: master data + records ----------
def err(msg, code=400): return jsonify(error=msg), code

def parse_date(s):
    try: return dt.date.fromisoformat(s).isoformat()
    except (TypeError, ValueError): return None

@app.get("/api/meta")
def meta():
    return jsonify(
        activities=rows("SELECT scope, activity, MAX(unit) unit FROM emission_factors GROUP BY scope, activity ORDER BY scope, activity"),
        years=[r["y"] for r in rows("SELECT DISTINCT CAST(strftime('%Y',activity_date) AS INT) y FROM emission_records ORDER BY y")],
        metrics=rows("SELECT DISTINCT metric_name, unit, aggregation FROM business_metrics"))

@app.get("/api/factors")
def list_factors():
    sql, a = "SELECT * FROM emission_factors WHERE 1=1", []
    if request.args.get("scope"): sql += " AND scope=?"; a.append(request.args["scope"])
    if request.args.get("activity"): sql += " AND activity=?"; a.append(request.args["activity"])
    if request.args.get("date"): sql += " AND valid_from<=? AND (valid_to IS NULL OR valid_to>=?)"; a += [request.args["date"]] * 2
    return jsonify(rows(sql + " ORDER BY scope, activity, valid_from", a))

@app.post("/api/factors")
def add_factor():
    d = request.get_json(force=True, silent=True) or {}
    vf = parse_date(d.get("valid_from"))
    try: scope, val = int(d["scope"]), float(d["kgco2e_per_unit"]); act, unit = d["activity"], d["unit"]
    except (KeyError, ValueError, TypeError): return err("scope, activity, unit, kgco2e_per_unit, valid_from required")
    if scope not in (1, 2) or not vf or val < 0: return err("invalid scope / valid_from / factor")
    con = db()
    try:
        # close the currently open version the day before the new one starts
        con.execute("UPDATE emission_factors SET valid_to=date(?, '-1 day') WHERE scope=? AND activity=? "
                    "AND valid_to IS NULL AND valid_from<?", (vf, scope, act, vf))
        cur = con.execute("INSERT INTO emission_factors(scope,activity,unit,kgco2e_per_unit,source,valid_from,valid_to)"
                          " VALUES(?,?,?,?,?,?,?)", (scope, act, unit, val, d.get("source"), vf, parse_date(d.get("valid_to"))))
        con.commit()
    except sqlite3.IntegrityError: return err("a version with that valid_from already exists", 409)
    return jsonify(id=cur.lastrowid), 201

@app.post("/api/emissions")
def add_emission():
    d = request.get_json(force=True, silent=True) or {}
    date = parse_date(d.get("activity_date"))
    try: scope, qty = int(d["scope"]), float(d["quantity"])
    except (KeyError, ValueError, TypeError): return err("scope and numeric quantity required")
    if scope not in (1, 2): return err("only scope 1 and 2 are supported")
    if not d.get("activity") or not date or qty <= 0: return err("activity, activity_date (YYYY-MM-DD), quantity > 0 required")
    try:
        rid = create_record(db(), scope, d["activity"], qty, date, d.get("section"), d.get("plant"), d.get("notes"))
    except ValueError as e: return err(str(e), 422)
    db().commit()
    return jsonify(rows("SELECT * FROM emission_records WHERE id=?", (rid,))[0]), 201

@app.get("/api/emissions")
def list_emissions():
    sql, a = "SELECT * FROM emission_records WHERE 1=1", []
    for k in ("scope", "origin"):
        if request.args.get(k): sql += f" AND {k}=?"; a.append(request.args[k])
    if request.args.get("year"): sql += " AND strftime('%Y',activity_date)=?"; a.append(request.args["year"])
    return jsonify(rows(sql + " ORDER BY id DESC LIMIT ?", a + [min(int(request.args.get("limit", 20)), 500)]))

@app.post("/api/emissions/<int:rid>/override")
def override(rid):
    d = request.get_json(force=True, silent=True) or {}
    rec = db().execute("SELECT * FROM emission_records WHERE id=?", (rid,)).fetchone()
    if not rec: return err("record not found", 404)
    reason = (d.get("reason") or "").strip()
    try: new = float(d["value_kgco2e"]); assert new >= 0
    except (KeyError, ValueError, TypeError, AssertionError): return err("value_kgco2e (>= 0) required")
    if not reason: return err("a reason is mandatory for manual overrides")
    db().execute("INSERT INTO audit_log(record_id,field,old_value,new_value,reason,changed_by) VALUES(?,?,?,?,?,?)",
                 (rid, "final_kgco2e", rec["final_kgco2e"], new, reason, d.get("user", "anonymous")))
    db().execute("UPDATE emission_records SET final_kgco2e=?, is_overridden=1 WHERE id=?", (new, rid))
    db().commit()
    return jsonify(rows("SELECT * FROM emission_records WHERE id=?", (rid,))[0])

@app.get("/api/audit-log")
def audit():
    sql, a = "SELECT * FROM audit_log", []
    if request.args.get("record_id"): sql += " WHERE record_id=?"; a.append(request.args["record_id"])
    return jsonify(rows(sql + " ORDER BY id DESC LIMIT 100", a))

@app.route("/api/business-metrics", methods=["GET", "POST"])
def metrics():
    if request.method == "GET":
        return jsonify(rows("SELECT * FROM business_metrics ORDER BY date DESC, metric_name LIMIT 200"))
    d = request.get_json(force=True, silent=True) or {}
    date = parse_date(d.get("date"))
    try: val = float(d["value"])
    except (KeyError, ValueError, TypeError): return err("numeric value required")
    if not date or not d.get("metric_name"): return err("date (YYYY-MM-DD) and metric_name required")
    db().execute("INSERT INTO business_metrics(date,metric_name,value,unit,aggregation) VALUES(?,?,?,?,?) "
                 "ON CONFLICT(date,metric_name) DO UPDATE SET value=excluded.value",
                 (date, d["metric_name"], val, d.get("unit"), d.get("aggregation", "sum")))
    db().commit()
    return jsonify(ok=True), 201

# ---------- API: analytics (all use final_kgco2e, i.e. overrides included) ----------
def by_year(y):
    out = {"scope1": 0.0, "scope2": 0.0}
    for r in rows("SELECT scope, SUM(final_kgco2e) s FROM emission_records WHERE strftime('%Y',activity_date)=? GROUP BY scope", (str(y),)):
        out[f"scope{r['scope']}"] = r["s"]
    out["total"] = out["scope1"] + out["scope2"]
    return out

@app.get("/api/analytics/yoy")
def yoy():
    y = int(request.args.get("year", dt.date.today().year))
    cur, prev = by_year(y), by_year(y - 1)
    chg = (cur["total"] - prev["total"]) / prev["total"] * 100 if prev["total"] else None
    return jsonify(year=y, previous_year=y - 1, current=cur, previous=prev, total_change_pct=chg, unit="kgCO2e")

@app.get("/api/analytics/intensity")
def intensity():
    y = request.args.get("year")
    start = request.args.get("start") or (f"{y}-01-01" if y else None)
    end = request.args.get("end") or (f"{y}-12-31" if y else None)
    if not parse_date(start) or not parse_date(end): return err("give year, or start and end (YYYY-MM-DD)")
    name = request.args.get("metric", "Tons of Steel Produced")
    em = db().execute("SELECT COALESCE(SUM(final_kgco2e),0) FROM emission_records WHERE activity_date BETWEEN ? AND ?", (start, end)).fetchone()[0]
    m = db().execute("SELECT SUM(value) s, AVG(value) a, MAX(aggregation) agg, MAX(unit) u, COUNT(*) n FROM business_metrics "
                     "WHERE metric_name=? AND date BETWEEN ? AND ?", (name, start, end)).fetchone()
    denom = (m["a"] if m["agg"] == "avg" else m["s"]) if m["n"] else None
    return jsonify(start=start, end=end, metric=name, metric_unit=m["u"], metric_value=denom, total_kgco2e=em,
                   intensity=(em / denom if denom else None), intensity_unit=f"kgCO2e per {m['u'] or 'unit'}")

@app.get("/api/analytics/hotspots")
def hotspots():
    y, top = request.args.get("year"), int(request.args.get("top_n", 8))
    sql, a = "SELECT activity, scope, SUM(final_kgco2e) kg FROM emission_records WHERE 1=1", []
    if y: sql += " AND strftime('%Y',activity_date)=?"; a.append(y)
    if request.args.get("scope"): sql += " AND scope=?"; a.append(request.args["scope"])
    data = rows(sql + " GROUP BY activity, scope ORDER BY kg DESC", a)
    total = sum(r["kg"] for r in data) or 1
    head, tail = data[:top], data[top:]
    if tail: head.append(dict(activity="Other sources", scope=None, kg=sum(r["kg"] for r in tail)))
    for r in head: r["pct"] = r["kg"] / total * 100
    return jsonify(year=y, total_kgco2e=total, sources=head)

@app.get("/api/analytics/monthly")
def monthly():
    y = request.args.get("year", str(dt.date.today().year))
    out = {m: {"month": m, "scope1": 0.0, "scope2": 0.0} for m in range(1, 13)}
    for r in rows("SELECT CAST(strftime('%m',activity_date) AS INT) m, scope, SUM(final_kgco2e) kg FROM emission_records "
                  "WHERE strftime('%Y',activity_date)=? GROUP BY m, scope", (y,)):
        out[r["m"]][f"scope{r['scope']}"] = r["kg"]
    for v in out.values(): v["total"] = v["scope1"] + v["scope2"]
    return jsonify(year=int(y), months=list(out.values()))

@app.get("/api/health")
def health(): return jsonify(status="ok")

@app.get("/")
def index(): return send_from_directory("static", "index.html")

init_db()
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=False)
